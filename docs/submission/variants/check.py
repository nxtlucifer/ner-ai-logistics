"""Compliance gate for every variant deck in this folder.

    python3 docs/submission/variants/check.py

Fails loudly on: wrong slide count, a shape off the canvas, non-wrapping text
wider than its box, a missing required string, or a banned claim string.
"""
from __future__ import annotations

import glob
import sys
from pathlib import Path

from PIL import ImageFont
from pptx import Presentation

HERE = Path(__file__).resolve().parent
FONTS = {
    ("Calibri", False): "/usr/share/fonts/truetype/crosextra/Carlito-Regular.ttf",
    ("Calibri", True): "/usr/share/fonts/truetype/crosextra/Carlito-Bold.ttf",
    ("Arial", False): "/usr/share/fonts/truetype/liberation/LiberationSans-Regular.ttf",
    ("Arial", True): "/usr/share/fonts/truetype/liberation/LiberationSans-Bold.ttf",
    ("Courier New", False): "/usr/share/fonts/truetype/liberation/LiberationMono-Regular.ttf",
    ("Courier New", True): "/usr/share/fonts/truetype/liberation/LiberationMono-Bold.ttf",
}
REQUIRED = ["SIH26002", "NER-AI LOGISTICS", "17", "Smart Automation", "Software",
            "UNKNOWN", "NOT DEPLOYED", "95 %", "FPR 50 %", "Precision 0.32", "F1 0.48"]
BANNED = ["98 %", "98%", "guaranteed safe", "AI detected landslide",
          "live landslide detection", "Google-level traffic", "50 % accuracy"]
_cache = {}


def width(txt, size, bold, name):
    key = (name if (name, bool(bold)) in FONTS else "Calibri", bool(bold), size)
    if key not in _cache:
        _cache[key] = ImageFont.truetype(FONTS[(key[0], key[1])], int(size * 4))
    return _cache[key].getlength(txt) / 4 / 72.0


def check(path: Path):
    prs = Presentation(str(path))
    sw, sh = prs.slide_width / 914400, prs.slide_height / 914400
    problems, seen = [], []
    if len(prs.slides) != 6:
        problems.append(f"slide count {len(prs.slides)} (want 6)")
    for i, slide in enumerate(prs.slides, 1):
        for shape in slide.shapes:
            try:
                L, T = shape.left / 914400, shape.top / 914400
                W, H = shape.width / 914400, shape.height / 914400
            except TypeError:
                continue
            if L < -0.7 or T < -0.7 or L + W > sw + 0.06 or T + H > sh + 0.06:
                problems.append(f"s{i} off-canvas: {shape.name} -> {L + W:.2f},{T + H:.2f}")
            if not shape.has_text_frame:
                continue
            tf = shape.text_frame
            for para in tf.paragraphs:
                runs = [r for r in para.runs if r.text]
                if not runs:
                    continue
                seen.append("".join(r.text for r in runs))
                if tf.word_wrap:
                    continue
                need = sum(width(r.text, r.font.size.pt if r.font.size else 12,
                                 r.font.bold, r.font.name or "Calibri") for r in runs)
                if need > W - 0.06:
                    problems.append(f"s{i} overflow: {shape.name} needs {need:.2f}in in {W:.2f}in "
                                    f"-- {''.join(r.text for r in runs)[:44]!r}")
    blob = "\n".join(seen)
    for token in REQUIRED:
        if token not in blob:
            problems.append(f"missing required string {token!r}")
    for token in BANNED:
        if token.lower() in blob.lower():
            problems.append(f"BANNED string present {token!r}")
    return problems


def main():
    bad = 0
    for f in sorted(glob.glob(str(HERE / "*.pptx"))):
        probs = check(Path(f))
        name = Path(f).stem.replace("RASTA_AI_SIH26002_TEAM17_", "")
        if probs:
            bad += 1
            print(f"FAIL {name}")
            for p in probs:
                print("      " + p)
        else:
            print(f"ok   {name}")
    print(("\n%d deck(s) need work" % bad) if bad else "\nall decks pass")
    return 1 if bad else 0


if __name__ == "__main__":
    sys.exit(main())
