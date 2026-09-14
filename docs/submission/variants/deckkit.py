"""Shared kit for the RASTA AI deck variants (SIH26002, Team 17).

Every variant is six slides on the supplied SIH2026 template, in the six
required categories, and may only state facts from docs/PPT_SOURCE_OF_TRUTH.md.
The facts live in FACTS below so a variant can restyle the deck but cannot
invent a number. Pictures are real screens from docs/submission/screenshots/.

Fonts are restricted to Calibri / Arial / Times New Roman / Courier New: those
four are the only ones with metric-compatible substitutes on the Linux renderer
(Carlito / Liberation Sans / Serif / Mono), so a variant looks the same in
PowerPoint and in the exported PDF.
"""
from __future__ import annotations

import os
from pathlib import Path

from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[2]
SHOTS = ROOT / "docs" / "submission" / "screenshots"
_WIN = Path(r"D:\SIH2026 PPT Format.pptx")
TEMPLATE = (Path(os.environ["SIH_TEMPLATE"]) if os.environ.get("SIH_TEMPLATE")
            else (_WIN if _WIN.exists() else HERE.parent / "template.pptx"))
CROPS = Path(os.environ.get("TEMP", "/tmp")) / "sih-variants"
CROPS.mkdir(parents=True, exist_ok=True)

SW, SH = 13.3333, 7.5          # slide, inches
FOOT = 6.95                    # the template's blue footer band starts here
BRAND = ROOT / "driver-app" / "assets" / "brand-mark.png"


def rgb(h):
    return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))


WHITE = rgb("FFFFFF")


# ------------------------------------------------------------------ content
FACTS = dict(
    name="RASTA AI",
    tagline="Route Intelligence for Essential-Supply Logistics in North Eastern India",
    ps_id="SIH26002",
    ps_title="AI-Based Smart Logistics and Accessibility Intelligence Platform for North Eastern Region (NER)",
    theme="Smart Automation",
    category="Software",
    team_id="17  (internal group no.)",
    team="NER-AI LOGISTICS",
    ministry="MDoNER",
    headline="SHORTEST IS NOT ALWAYS USABLE.",
    question="\u201cCan this essential-supply truck reliably use this corridor now?\u201d",
    proof_line="End-to-end working prototype \u00b7 physical-phone validated \u00b7 live demo corridor Guwahati \u2192 Shillong",
    corridor="Guwahati \u2192 Shillong \u00b7 \u2248 98.8 km real route",
    problem="Monsoon \u00b7 landslides \u00b7 floods \u00b7 official warnings \u00b7 single-lane hill roads \u00b7 no signal.",
    problem_kick="A shortest-path router still sends the truck.",
    normal_router=["Origin", "Shortest / fastest road", "Truck"],
    normal_gap=["No terrain, weather or hazard check.", "No governance. No answer when data is missing."],
    evidence_chips=["Terrain", "Weather", "Landslide exposure", "Flood context",
                    "Official warnings", "Fleet traffic", "Evidence freshness"],
    decisions=[("CONTINUE", "ok"), ("CAUTION", "warn"), ("HOLD", "danger"), ("REROUTE", "ink")],
    unknown_rule="Missing or stale evidence is marked UNKNOWN \u2014 never safe \u2014 and goes to an authorised reviewer.",
    chain=["Verified evidence", "Deterministic safety policy", "Manager governance", "Driver action"],
    sources=[("WEATHER", "Open-Meteo \u00b7 MET Norway"),
             ("TERRAIN", "Copernicus DEM \u00b7 OpenTopoData"),
             ("WARNINGS", "NDMA SACHET (CAP)"),
             ("FLOOD", "GloFAS"),
             ("LANDSLIDE HISTORY", "NASA GLC 2007\u201317"),
             ("TRAFFIC", "fleet GPS fixes")],
    deciders=[("DETERMINISTIC SAFETY ENGINE", "makes every operational decision", "accent"),
              ("GEMINI / OPENROUTER (LLM)", "explains the decision in the driver's language \u2014 never decides", "ok"),
              ("EXPERIMENTAL ML", "research channel only; held outside routing by the safety gate", "muted")],
    stack="FastAPI \u00b7 PostgreSQL / PostGIS (Supabase) \u00b7 React \u00b7 React Native / Expo \u00b7 OSRM \u00b7 OpenStreetMap \u00b7 Gemini",
    clients="Manager console (React + TS) \u00b7 Driver app (React Native / Expo, Android) \u00b7 one login, server-decided role, offline package",
    run_line="Run end to end on a physical Android phone, hosted backend, real road.",
    screens_note="APK 1.0.18. One login, server-decided role. Every language shows its status \u2014 Verified, Draft or English fallback.",
    loop_note="Re-scored every 60 s on the road ahead. A reroute is proposed to the manager, never applied silently.",
    lifecycle=["PLAN", "REVIEW", "DISPATCH", "ACCEPT", "TRUCK VERIFY", "NAVIGATE",
               "OFF-ROUTE", "MANAGER REROUTE APPROVAL", "DELIVERY"],
    proof=[("12/12", "physical judge-flow steps"), ("9/9", "physical role-flow steps")],
    tests=[("1138+", "backend"), ("621+", "driver"), ("170+", "manager")],
    limits=("Cold start \u2192 warm /health first  \u00b7  rate limits \u2192 cached evidence + explicit UNKNOWN  \u00b7  "
            "push needs Firebase \u2192 in-app alerts work  \u00b7  experimental ML \u2192 outside routing"),
    supplies="FOOD  \u00b7  MEDICINE  \u00b7  FUEL  \u00b7  RELIEF SUPPLIES",
    who=[("FLEET MANAGERS", "Usability and uncertainty visible before dispatch. Auditable."),
         ("DRIVERS", "Danger context, emergency support, 23 languages."),
         ("REMOTE COMMUNITIES", "Resilient access to food, medicine, fuel, relief."),
         ("GOVERNMENT / OPS", "Auditable evidence and disruption per corridor.")],
    caps=[("23", "selectable languages, explicit fallback status"),
          ("11", "route factors in every decision"),
          ("60 s", "route-ahead monitor while driving"),
          ("12/12", "judge-flow steps certified on a physical phone"),
          ("100 %", "of reroutes reviewed by a human"),
          ("0", "unknown factors counted as safe")],
    groups=[("ROUTING / PLACES", "OSRM \u00b7 OpenStreetMap \u00b7 Nominatim",
             "Road geometry, alternatives, turn steps."),
            ("WEATHER / TERRAIN", "Open-Meteo \u00b7 MET Norway \u00b7 Copernicus DEM \u00b7 OpenTopoData",
             "Forecast along the road; climb, gradient, steep segments."),
            ("OFFICIAL / FLOOD", "NDMA SACHET (CAP) \u00b7 Copernicus GloFAS",
             "Official warnings with freshness; discharge vs 30-day mean."),
            ("LANDSLIDE HISTORY", "NASA Global Landslide Catalog",
             "Historical exposure, 2007\u20132017, \u22645 km accuracy."),
            ("AI ASSISTANCE", "Google Gemini",
             "Explains in the driver's language. Never safety authority."),
            ("PROBLEM STATEMENT", "MDoNER \u00b7 SIH26002",
             "Smart logistics and accessibility intelligence for the NER.")],
    ml_recall="95 %",
    ml_recall_what="RECALL",
    ml_recall_sub="NER geographic holdout \u00b7 154 events \u00b7 no NER row in training",
    ml_gate_small="SAFETY GATE",
    ml_gate_big="NOT DEPLOYED",
    ml_fpr="FPR 50 %",
    ml_fpr_why="too high for production deployment",
    ml_flow=["TRAIN", "HELD-OUT NER TEST", "USEFUL SIGNAL FOUND",
             "FALSE POSITIVES TOO HIGH", "SAFETY GATE", "NOT DEPLOYED"],
    ml_caption_lead="VALIDATION BEFORE AUTOMATION.",
    ml_caption="Recall was not enough. Held-out testing exposed the false-positive rate. The gate keeps it out of routing.",
    ml_metrics=("Precision 0.32 \u00b7 F1 0.48 \u00b7 PR-AUC 0.55 \u00b7 ROC-AUC 0.83 \u00b7 Brier 0.12  \u2014  "
                "calibrated logistic regression, NASA GLC + ERA5-Land rain + DEM slope"),
    refs=("gpm.nasa.gov/landslides  \u00b7  sachet.ndma.gov.in  \u00b7  global-flood.emergency.copernicus.eu  \u00b7  "
          "open-meteo.com  \u00b7  api.met.no  \u00b7  opentopodata.org  \u00b7  project-osrm.org  \u00b7  openstreetmap.org  \u00b7  "
          "github.com/nxtlucifer/ner-ai-logistics"),
    # --- what makes it different (SIH judges ask for USP explicitly)
    usp=[("CORRIDOR USABILITY, NOT SHORTEST PATH",
          "11 factors scored on the selected road before dispatch"),
         ("UNKNOWN IS NEVER SAFE",
          "missing or stale evidence blocks selection until a reviewer authorises"),
         ("HUMAN-GOVERNED REROUTE",
          "a real off-route triggers a real alternative; a manager accepts before the phone follows"),
         ("ONE APK, 23 LANGUAGES, OFFLINE",
          "server-decided role; whole-trip package cached on the phone")],
    # --- the four things the feasibility slide must answer
    feasibility=[("TECHNOLOGY",
                  "FastAPI + PostGIS, two clients on one REST API. Every part is built and running."),
                 ("IMPLEMENTATION",
                  "Hosted backend, manager console, Android APK 1.0.18. Certified on a physical phone."),
                 ("COST",
                  "No licensed map, traffic or weather feed \u2014 every source is public or open."),
                 ("SCALABILITY",
                  "A new source is a new provider adapter behind the same policy. Fleet traffic improves per truck.")],
    risks=[("Provider rate limit or outage", "Cached evidence; the factor shows UNKNOWN, never assumed safe"),
           ("Cold start on free hosting", "Warm /health before the slot; state is server-side"),
           ("No phone signal on the corridor", "Whole-trip offline package; navigation continues, labelled OFFLINE"),
           ("Experimental model unreliable", "Safety gate holds it outside routing until it passes")],
    evidence_facts=[("98.8 km", "real road on the demo corridor"),
                    ("22", "landslides within 5 km"),
                    ("1975 m", "climb, steepest 10.3 %"),
                    ("154", "NER events, NASA GLC")],
    titles=["", "RASTA AI \u2014 ROUTE INTELLIGENCE", "TECHNICAL APPROACH",
            "FEASIBILITY AND VIABILITY", "IMPACT AND BENEFITS", "RESEARCH  AND REFERENCES"],
)

SHOT_SPECS = {
    "mgr":    ("02-trip-route.png", (1370, 200, 2470, 1800)),
    "evid":   ("03-route-evidence.png", (1400, 70, 2460, 1720)),
    "nav":    ("07-navigation.png", None),
    "truck":  ("06-truck-verification.png", None),
    "lang":   ("11-language-selector.png", (0, 1000, 1264, 2780)),
    "mobile": ("10-manager-mobile.png", (0, 100, 1264, 2780)),
    "review": ("09b-reroute-manager-review.png", (1340, 150, 2470, 1760)),
    "diag":   ("13-diagnostics.png", (1360, 150, 2640, 1760)),
    "ai":     ("08-route-ai.png", None),
    "trip":   ("06b-trip-after-verification.png", (0, 120, 1264, 2500)),
}


def shots():
    """Crop every screen once; returns {key: (path, (w, h))}."""
    out = {}
    for key, (name, box) in SHOT_SPECS.items():
        im = Image.open(SHOTS / name).convert("RGB")
        if box:
            im = im.crop(box)
        dst = CROPS / f"{key}.png"
        if not dst.exists():
            im.save(dst, "PNG", optimize=True)
        out[key] = (dst, im.size)
    return out


# ------------------------------------------------------------------- theme
class Theme:
    def __init__(self, **kw):
        self.font = kw.get("font", "Calibri")
        self.display = kw.get("display", self.font)
        self.mono = kw.get("mono", "Courier New")
        self.ink = rgb(kw.get("ink", "101820"))
        self.accent = rgb(kw.get("accent", "2563EB"))
        self.accent2 = rgb(kw.get("accent2", kw.get("accent", "2563EB")))
        self.ok = rgb(kw.get("ok", "059669"))
        self.warn = rgb(kw.get("warn", "D97706"))
        self.danger = rgb(kw.get("danger", "DC2626"))
        self.grey = rgb(kw.get("grey", "6B7280"))
        self.muted = rgb(kw.get("muted", "9CA3AF"))
        self.line = rgb(kw.get("line", "E5E7EB"))
        self.bg = rgb(kw.get("bg", "F7F8FA"))
        self.paper = rgb(kw.get("paper", "FFFFFF"))
        self.on_dark = rgb(kw.get("on_dark", "FFFFFF"))
        self.dark = kw.get("dark", False)
        self.radius = kw.get("radius", 0.1)
        self.border = kw.get("border", 1.25)
        self.slug = kw["slug"]
        self.label = kw["label"]
        self.note = kw["note"]

    def c(self, key):
        return {"ink": self.ink, "accent": self.accent, "accent2": self.accent2, "ok": self.ok,
                "warn": self.warn, "danger": self.danger, "grey": self.grey, "muted": self.muted,
                "line": self.line, "bg": self.bg, "paper": self.paper, "white": WHITE,
                "on_dark": self.on_dark}[key]


# ---------------------------------------------------------------- drawing
def _no_shadow(shape):
    spPr = shape._element.spPr
    if spPr.find(qn("a:effectLst")) is None:
        etree.SubElement(spPr, qn("a:effectLst"))


def rect(s, x, y, w, h, fill=None, line=None, rounded=False, line_w=0.75, radius=None):
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE
    sh = s.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    if rounded:
        sh.adjustments[0] = 0.12 if radius is None else min(0.5, radius / max(0.01, min(w, h)))
    if fill is None:
        sh.fill.background()
    else:
        sh.fill.solid(); sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line; sh.line.width = Pt(line_w)
    _no_shadow(sh)
    sh.text_frame.text = ""
    return sh


def text(s, x, y, w, h, paras, t=None, size=12, color=None, bold=False, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, spacing=1.05, space_after=0, font=None, margin=0.02):
    """paras: list of paragraphs; each is a str or a list of (text, opts) runs."""
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(margin)
    tf.margin_top = tf.margin_bottom = Inches(margin)
    base_font = font or (t.font if t else "Calibri")
    base_color = color if color is not None else (t.ink if t else rgb("101820"))
    first = True
    for para in paras:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = align
        p.line_spacing = spacing
        p.space_after = Pt(space_after)
        for txt, o in ([(para, {})] if isinstance(para, str) else para):
            r = p.add_run(); r.text = txt
            f = r.font
            f.name = o.get("font", base_font)
            f.size = Pt(o.get("size", size))
            f.bold = o.get("bold", bold)
            f.color.rgb = o.get("color", base_color)
            if o.get("italic"):
                f.italic = True
    return tb


def pill(s, x, y, w, h, label, t, fill=None, color=None, size=10.5, line=None, bold=True,
         rounded=True, radius=None):
    sh = rect(s, x, y, w, h, fill=fill, line=line, rounded=rounded,
              radius=(h / 2 if radius is None and rounded else radius))
    tf = sh.text_frame
    tf.margin_left = tf.margin_right = Inches(0.06)
    tf.margin_top = tf.margin_bottom = Inches(0.0)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.word_wrap = True
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = label
    r.font.name = t.font; r.font.size = Pt(size); r.font.bold = bold
    r.font.color.rgb = color if color is not None else WHITE
    return sh


def arrow(s, x1, y1, x2, y2, color, width=1.25, head=True):
    c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = color; c.line.width = Pt(width)
    if head:
        ln = c.line._get_or_add_ln()
        e = etree.SubElement(ln, qn("a:tailEnd")); e.set("type", "triangle"); e.set("w", "med"); e.set("len", "med")
    return c


def hline(s, x, y, w, color, width=0.75):
    return arrow(s, x, y, x + w, y, color, width=width, head=False)


def picture(s, spec, x, y, w=None, h=None, frame=None, t=None, shadow=False):
    path, size = spec
    if w and not h:
        h = w * size[1] / size[0]
    if h and not w:
        w = h * size[0] / size[1]
    if frame is not None:
        rect(s, x - 0.035, y - 0.035, w + 0.07, h + 0.07, fill=WHITE, line=frame, rounded=True, radius=0.06)
    s.shapes.add_picture(str(path), Inches(x), Inches(y), Inches(w), Inches(h))
    return w, h


def caption(s, x, y, w, label, t, fill=None, color=None, size=9):
    return pill(s, x, y, w, 0.28, label, t, fill=fill or t.ink, color=color or WHITE, size=size)


def panel(s, x, y, w, h, t, title=None, sub=None, fill=None, line=None, line_w=None,
          title_size=11.5, title_color=None, align=PP_ALIGN.LEFT, pad=0.16):
    rect(s, x, y, w, h, fill=(t.paper if fill is None else fill),
         line=(t.accent if line is None else line), rounded=True,
         line_w=(t.border if line_w is None else line_w), radius=t.radius)
    if not title:
        return y + pad
    runs = [(title, {"bold": True, "size": title_size, "color": title_color or t.accent})]
    if sub:
        runs.append(("   " + sub, {"size": title_size - 1.5, "color": t.grey}))
    text(s, x + pad, y + 0.08, w - 2 * pad, 0.3, [runs], t=t, align=align)
    return y + 0.42


def stat(s, x, y, w, number, label, t, num_size=20, lab_size=11, num_color=None,
         lab_color=None, align=PP_ALIGN.LEFT, gap=1.05):
    text(s, x, y, gap - 0.08, 0.4, [[(number, {"bold": True, "size": num_size,
         "color": num_color or t.accent})]], t=t, anchor=MSO_ANCHOR.MIDDLE, align=align)
    text(s, x + gap, y, w - gap, 0.4, [[(label, {"size": lab_size, "color": lab_color or t.ink})]],
         t=t, anchor=MSO_ANCHOR.MIDDLE)


def chain_h(s, x, y, items, t, h=0.34, gap=0.11, size=9, fills=None, colors=None, lines=None):
    """items: list of (label, width). fills/colors/lines: per-index overrides or None."""
    for i, (lab, w) in enumerate(items):
        fill = (fills or {}).get(i, t.paper)
        col = (colors or {}).get(i, t.ink)
        ln = (lines or {}).get(i, t.ink)
        pill(s, x, y, w, h, lab, t, fill=fill, color=col, line=ln, size=size)
        if i < len(items) - 1:
            arrow(s, x + w + 0.02, y + h / 2, x + w + gap - 0.02, y + h / 2, t.ink)
        x += w + gap
    return x


def chain_v(s, x, w, y, items, t, h=0.3, gap=0.12, size=10.5):
    """items: list of (label, fill, color, line)."""
    for i, (lab, fill, col, ln) in enumerate(items):
        pill(s, x, y, w, h, lab, t, fill=fill, color=col, line=ln, size=size)
        if i < len(items) - 1:
            arrow(s, x + w / 2, y + h + 0.01, x + w / 2, y + h + gap - 0.01, t.ink)
        y += h + gap
    return y


# ------------------------------------------------------------- template ops
def strip(slide, names):
    for sh in list(slide.shapes):
        if sh.name in names:
            sh._element.getparent().remove(sh._element)


def instruction_box(slide):
    strip(slide, {"TextBox 8"})


def team_oval(slide, t, label=None):
    for sh in slide.shapes:
        if sh.name.startswith("Oval"):
            sh.width = Inches(1.7); sh.height = Inches(0.8)
            sh.left = Inches(0.3); sh.top = Inches(0.3)
            tf = sh.text_frame; tf.word_wrap = True
            tf.margin_left = tf.margin_right = Inches(0.04)
            tf.vertical_anchor = MSO_ANCHOR.MIDDLE
            for p in list(tf.paragraphs)[1:]:
                p._p.getparent().remove(p._p)
            p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
            for r in list(p.runs):
                r._r.getparent().remove(r._r)
            r = p.add_run(); r.text = label or FACTS["team"]
            r.font.name = t.font; r.font.size = Pt(10); r.font.bold = True; r.font.color.rgb = t.ink
            return sh


def set_title(slide, s, size=None, left=None, width=None, color=None, font=None):
    """Retitle a content slide.

    The template's own title boxes sit at top=-0.05 with 36 pt type, so any
    smaller size rides up off the top edge of the slide. Whenever the size is
    changed the box is re-seated at a known-good geometry and anchored middle.
    One of them also carries a stray line break before its run, which would
    push the new title onto a second line; it is dropped.
    """
    for sh in slide.shapes:
        if not (sh.is_placeholder and sh.placeholder_format.type in (1, 3)):
            continue
        tf = sh.text_frame
        p = tf.paragraphs[0]
        for br in p._p.findall(qn("a:br")):
            p._p.remove(br)
        for r in list(p.runs)[1:]:
            r._r.getparent().remove(r._r)
        if p.runs:
            run = p.runs[0]
            run.text = s
            if size:
                run.font.size = Pt(size)
                p._p.get_or_add_endParaRPr().set("sz", str(int(size * 100)))
            if color is not None:
                run.font.color.rgb = color
            if font:
                run.font.name = font
        else:
            tf.text = s
        if size:
            sh.top = Inches(0.0)
            sh.height = Inches(1.16)
            tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        if left is not None:
            sh.left = Inches(left)
        if width is not None:
            sh.width = Inches(width)
        return sh


def title_slide_prep(slide, title_size=34, width=10.15):
    """Clear the template's own title-page furniture and stop the heading short
    of the SIH logo, which it otherwise prints on top of."""
    for sh in list(slide.shapes):
        if sh.name in ("Freeform: Shape 26", "Picture 4", "Subtitle 3", "TextBox 9"):
            sh._element.getparent().remove(sh._element)
        elif sh.name == "Title 7":
            sh.width = Inches(width)
            for para in sh.text_frame.paragraphs:
                for r in para.runs:
                    r.font.size = Pt(title_size)
                para._p.get_or_add_endParaRPr().set("sz", str(int(title_size * 100)))


def content_slide(slide, t, title, size=28, left=2.1, width=8.5):
    instruction_box(slide)
    team_oval(slide, t)
    set_title(slide, title, size=size, left=left, width=width)


def dark_ground(slide, t, color):
    """Paint the whole slide, behind everything the template already draws."""
    sh = rect(slide, -0.02, -0.02, SW + 0.04, FOOT + 0.02, fill=color)
    slide.shapes._spTree.remove(sh._element)
    slide.shapes._spTree.insert(2, sh._element)
    return sh


def finish(prs, out: Path):
    lst = prs.slides._sldIdLst
    last = list(lst)[6]
    prs.part.drop_rel(last.get(qn("r:id")))
    lst.remove(last)
    cp = prs.core_properties
    cp.title = f"{FACTS['name']} \u2014 {FACTS['ps_id']} \u2014 Team {FACTS['team']} ({FACTS['team_id'].split()[0]})"
    cp.subject = FACTS["ps_title"]
    cp.author = FACTS["team"]
    cp.last_modified_by = FACTS["team"]
    cp.category = f"Smart India Hackathon 2026 \u2014 {FACTS['theme']} \u2014 {FACTS['category']}"
    cp.keywords = "SIH26002, SIH 2026, RASTA AI, NER logistics, route intelligence"
    out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out))
    return out


def new_deck():
    return Presentation(str(TEMPLATE))
