"""Build the Day 1 Task 1 document as DOCX (python-docx) and PDF (reportlab)
from RASTA_AI_SIH2026_DAY1_TASK1_SYSTEM_DESIGN.md. Mermaid blocks are replaced
by the PNGs in diagrams/ (rendered by render_diagrams.mjs); wide diagrams and
tables go on landscape pages. No external converter is needed.

    python docs/submission/day1/build_docs.py
"""
from __future__ import annotations

import re
import sys
from dataclasses import dataclass, field
from pathlib import Path

HERE = Path(__file__).resolve().parent
MD = HERE / "RASTA_AI_SIH2026_DAY1_TASK1_SYSTEM_DESIGN.md"
DIAGRAMS = HERE / "diagrams"
OUT_DOCX = HERE / "RASTA_AI_SIH2026_DAY1_TASK1_SYSTEM_DESIGN.docx"
OUT_PDF = HERE / "RASTA_AI_SIH2026_DAY1_TASK1_SYSTEM_DESIGN.pdf"

DIAGRAM_CAPTIONS = {
    1: "Figure 1 — High-level architecture (rendered from the Mermaid source in the Markdown edition)",
    2: "Figure 2 — Data flow diagram, Level 0 (context)",
    3: "Figure 3 — Data flow diagram, Level 1 (processes, stores, external providers)",
}
LANDSCAPE_DIAGRAMS = {3}
LANDSCAPE_MIN_COLUMNS = 5  # tables with this many columns or more go on landscape pages

# --------------------------------------------------------------------------- parse


@dataclass
class Block:
    kind: str  # heading | para | bullets | numbers | table | code | image | quote | rule
    text: str = ""
    level: int = 0
    items: list[str] = field(default_factory=list)
    rows: list[list[str]] = field(default_factory=list)
    lang: str = ""
    index: int = 0


def parse(md: str) -> list[Block]:
    lines = md.split("\n")
    blocks: list[Block] = []
    i = 0
    diagram = 0
    para: list[str] = []

    def flush_para() -> None:
        nonlocal para
        if para:
            blocks.append(Block("para", " ".join(s.strip() for s in para)))
            para = []

    while i < len(lines):
        line = lines[i]
        if line.startswith("```"):
            flush_para()
            lang = line[3:].strip()
            body: list[str] = []
            i += 1
            while i < len(lines) and not lines[i].startswith("```"):
                body.append(lines[i])
                i += 1
            i += 1
            if lang == "mermaid":
                diagram += 1
                blocks.append(Block("image", index=diagram))
            else:
                blocks.append(Block("code", "\n".join(body), lang=lang))
            continue
        m = re.match(r"^(#{1,4})\s+(.*)$", line)
        if m:
            flush_para()
            blocks.append(Block("heading", m.group(2).strip(), level=len(m.group(1))))
            i += 1
            continue
        if line.strip() == "---":
            flush_para()
            blocks.append(Block("rule"))
            i += 1
            continue
        if line.startswith("|"):
            flush_para()
            rows: list[list[str]] = []
            while i < len(lines) and lines[i].startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-{3,}:?", c) for c in cells):
                    rows.append(cells)
                i += 1
            blocks.append(Block("table", rows=rows))
            continue
        if re.match(r"^\s*[-*]\s+", line):
            flush_para()
            items: list[str] = []
            while i < len(lines) and re.match(r"^\s*[-*]\s+", lines[i]):
                items.append(re.sub(r"^\s*[-*]\s+", "", lines[i]).strip())
                i += 1
            blocks.append(Block("bullets", items=items))
            continue
        if re.match(r"^\s*\d+\.\s+", line):
            flush_para()
            items = []
            while i < len(lines) and re.match(r"^\s*\d+\.\s+", lines[i]):
                items.append(re.sub(r"^\s*\d+\.\s+", "", lines[i]).strip())
                i += 1
            blocks.append(Block("numbers", items=items))
            continue
        if line.startswith(">"):
            flush_para()
            q: list[str] = []
            while i < len(lines) and lines[i].startswith(">"):
                q.append(lines[i].lstrip("> ").strip())
                i += 1
            blocks.append(Block("quote", " ".join(q)))
            continue
        if not line.strip():
            flush_para()
            i += 1
            continue
        para.append(line)
        i += 1
    flush_para()
    return blocks


INLINE = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`|\*[^*\s][^*]*\*)")


def inline_runs(text: str) -> list[tuple[str, str]]:
    """Split markdown inline text into (style, text) runs: plain | bold | code | italic."""
    runs: list[tuple[str, str]] = []
    pos = 0
    for m in INLINE.finditer(text):
        if m.start() > pos:
            runs.append(("plain", text[pos:m.start()]))
        tok = m.group(0)
        if tok.startswith("**"):
            runs.append(("bold", tok[2:-2]))
        elif tok.startswith("`"):
            runs.append(("code", tok[1:-1]))
        else:
            runs.append(("italic", tok[1:-1]))
        pos = m.end()
    if pos < len(text):
        runs.append(("plain", text[pos:]))
    return runs


def wants_landscape(b: Block) -> bool:
    return (b.kind == "image" and b.index in LANDSCAPE_DIAGRAMS) or (
        b.kind == "table" and bool(b.rows) and len(b.rows[0]) >= LANDSCAPE_MIN_COLUMNS
    )


def landscape_plan(body: list[Block]) -> list[bool]:
    """Per block: should it sit on a landscape page? A heading (or a short lead
    paragraph) directly before a landscape table/diagram goes with it, so no
    page is left holding a heading alone."""
    plan = [wants_landscape(b) for b in body]
    for i in range(len(body) - 1, -1, -1):
        if body[i].kind in ("heading", "para") and not plan[i]:
            j = i + 1
            while j < len(body) and body[j].kind == "rule":
                j += 1
            if j < len(body) and plan[j] and (body[i].kind == "heading" or body[j].kind == "table"):
                plan[i] = True
    return plan


def split_front_matter(blocks: list[Block]) -> tuple[list[Block], list[Block]]:
    """Everything before '## Table of contents' is the title page; the markdown
    TOC itself is dropped (both editions generate their own)."""
    for n, b in enumerate(blocks):
        if b.kind == "heading" and b.text.lower().startswith("table of contents"):
            front = blocks[:n]
            rest = blocks[n + 1 :]
            # skip the TOC's own numbered list and the rule after it
            while rest and rest[0].kind in ("numbers", "rule"):
                rest.pop(0)
            return front, rest
    return [], blocks


# --------------------------------------------------------------------------- docx

def build_docx(front: list[Block], body: list[Block]) -> None:
    from docx import Document
    from docx.enum.section import WD_ORIENT
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor

    doc = Document()
    st = doc.styles
    st["Normal"].font.name = "Calibri"
    st["Normal"].font.size = Pt(10.5)
    for name, size in (("Heading 1", 20), ("Heading 2", 15), ("Heading 3", 12.5), ("Heading 4", 11)):
        st[name].font.name = "Calibri"
        st[name].font.size = Pt(size)
        st[name].font.color.rgb = RGBColor(0x14, 0x3D, 0x2E)
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    for s in (sec,):
        s.left_margin = s.right_margin = Cm(2.0)
        s.top_margin = s.bottom_margin = Cm(2.0)

    def add_page_number(section) -> None:
        p = section.footer.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run("RASTA AI · SIH 2026 · Day 1 Task 1 · Page ")
        r.font.size = Pt(8)
        for tag, text in (("begin", None), (None, "PAGE"), ("end", None)):
            run = p.add_run()
            run.font.size = Pt(8)
            if tag:
                fc = OxmlElement("w:fldChar")
                fc.set(qn("w:fldCharType"), tag)
                run._r.append(fc)
            else:
                it = OxmlElement("w:instrText")
                it.set(qn("xml:space"), "preserve")
                it.text = text
                run._r.append(it)

    add_page_number(sec)

    def add_runs(par, text: str, size: float | None = None) -> None:
        for style, chunk in inline_runs(text):
            r = par.add_run(chunk)
            if size:
                r.font.size = Pt(size)
            if style == "bold":
                r.bold = True
            elif style == "italic":
                r.italic = True
            elif style == "code":
                r.font.name = "Consolas"
                r._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
                r.font.size = Pt((size or 10.5) - 1)

    # ---- title page
    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    t.paragraph_format.space_before = Pt(140)
    r = t.add_run("RASTA AI")
    r.bold = True
    r.font.size = Pt(40)
    r.font.color.rgb = RGBColor(0x14, 0x3D, 0x2E)
    for line, size in (("SIH 2026 — Day 1 Task 1", 20), ("Problem Understanding, Solution Planning & System Design", 14)):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        rr = p.add_run(line)
        rr.font.size = Pt(size)
    doc.add_paragraph()
    meta = next((b for b in front if b.kind == "table"), None)
    if meta:
        tbl = doc.add_table(rows=0, cols=2)
        tbl.style = "Table Grid"
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        for row in meta.rows:
            if not any(c for c in row):
                continue
            cells = tbl.add_row().cells
            add_runs(cells[0].paragraphs[0], row[0], 9.5)
            add_runs(cells[1].paragraphs[0], row[1] if len(row) > 1 else "", 9.5)
        for row in tbl.rows:
            row.cells[0].width = Cm(5.5)
            row.cells[1].width = Cm(11.5)
    note = next((b for b in front if b.kind == "para" and b.text.startswith("**How to read")), None)
    if note:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(18)
        add_runs(p, note.text, 9.5)
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    # ---- table of contents (field; Word fills it on open)
    h = doc.add_paragraph("Table of contents", style="Heading 1")
    p = doc.add_paragraph()
    run = p.add_run()
    for tag, text in (("begin", None), (None, 'TOC \\o "1-2" \\h \\z \\u'), ("separate", None), (None, None), ("end", None)):
        if tag:
            fc = OxmlElement("w:fldChar")
            fc.set(qn("w:fldCharType"), tag)
            run._r.append(fc)
        elif text:
            it = OxmlElement("w:instrText")
            it.set(qn("xml:space"), "preserve")
            it.text = text
            run._r.append(it)
        else:
            tx = OxmlElement("w:t")
            tx.text = "Right-click and choose Update Field (or press F9) to fill this table of contents."
            run._r.append(tx)
    settings = doc.settings.element
    uf = OxmlElement("w:updateFields")
    uf.set(qn("w:val"), "true")
    settings.append(uf)
    # a plain list too, so the TOC reads even before the field is updated
    for b in body:
        if b.kind == "heading" and b.level == 2:
            doc.add_paragraph(b.text, style="List Number")
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    # ---- body
    def new_section(landscape: bool):
        from docx.enum.section import WD_SECTION
        s = doc.add_section(WD_SECTION.NEW_PAGE)
        if landscape:
            s.orientation = WD_ORIENT.LANDSCAPE
            s.page_width, s.page_height = Cm(29.7), Cm(21.0)
        else:
            s.orientation = WD_ORIENT.PORTRAIT
            s.page_width, s.page_height = Cm(21.0), Cm(29.7)
        s.left_margin = s.right_margin = Cm(2.0)
        s.top_margin = s.bottom_margin = Cm(2.0)
        s.footer.is_linked_to_previous = True
        return s

    landscape = False
    first_h2 = True
    plan = landscape_plan(body)
    for b, want in zip(body, plan):
        just_broke = False
        if want and not landscape:
            new_section(True)
            landscape = True
            just_broke = True
        elif landscape and not want and b.kind in ("heading", "para", "bullets", "numbers", "code", "quote", "table", "image"):
            new_section(False)
            landscape = False
            just_broke = True
        if b.kind == "heading":
            if b.level == 2 and not first_h2 and not just_broke:
                doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
            first_h2 = False if b.level == 2 else first_h2
            style = {1: "Heading 1", 2: "Heading 1", 3: "Heading 2", 4: "Heading 3"}[b.level]
            hp = doc.add_paragraph(style=style)
            add_runs(hp, b.text)
        elif b.kind == "para":
            p = doc.add_paragraph()
            add_runs(p, b.text)
        elif b.kind == "quote":
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(1)
            p.paragraph_format.right_indent = Cm(1)
            add_runs(p, b.text)
            for r in p.runs:
                r.italic = True
        elif b.kind == "bullets":
            for it in b.items:
                p = doc.add_paragraph(style="List Bullet")
                add_runs(p, it)
        elif b.kind == "numbers":
            for it in b.items:
                p = doc.add_paragraph(style="List Number")
                add_runs(p, it)
        elif b.kind == "code":
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(6)
            pPr = p._p.get_or_add_pPr()
            shd = OxmlElement("w:shd")
            shd.set(qn("w:val"), "clear")
            shd.set(qn("w:fill"), "F3F5F4")
            pPr.append(shd)
            longest = max((len(l) for l in b.text.split("\n")), default=40)
            size = 8.5 if longest <= 78 else 7.5 if longest <= 95 else 6.5
            for n, l in enumerate(b.text.split("\n")):
                r = p.add_run(l)
                r.font.name = "Consolas"
                r._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
                r.font.size = Pt(size)
                if n < len(b.text.split("\n")) - 1:
                    r.add_break()
        elif b.kind == "table":
            cols = len(b.rows[0])
            tbl = doc.add_table(rows=0, cols=cols)
            tbl.style = "Table Grid"
            size = 9 if cols <= 3 else 8 if cols <= 5 else 7
            for n, row in enumerate(b.rows):
                cells = tbl.add_row().cells
                for c, text in enumerate(row[:cols]):
                    par = cells[c].paragraphs[0]
                    add_runs(par, text, size)
                    if n == 0:
                        for r in par.runs:
                            r.bold = True
                        tcPr = cells[c]._tc.get_or_add_tcPr()
                        shd = OxmlElement("w:shd")
                        shd.set(qn("w:val"), "clear")
                        shd.set(qn("w:fill"), "E3EBE6")
                        tcPr.append(shd)
            doc.add_paragraph()
        elif b.kind == "image":
            png = DIAGRAMS / f"diagram-{b.index}.png"
            if png.exists():
                width = Cm(25.5) if landscape else Cm(17.0)
                doc.add_picture(str(png), width=width)
                doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
                cap = doc.add_paragraph(DIAGRAM_CAPTIONS.get(b.index, f"Figure {b.index}"))
                cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for r in cap.runs:
                    r.italic = True
                    r.font.size = Pt(9)
        elif b.kind == "rule":
            pass
    doc.save(OUT_DOCX)
    print("docx:", OUT_DOCX.name, OUT_DOCX.stat().st_size // 1024, "KB")


# --------------------------------------------------------------------------- pdf

def build_pdf(front: list[Block], body: list[Block]) -> None:
    from fontTools.ttLib import TTFont as FTFont
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER
    from reportlab.lib.pagesizes import A4, landscape as LS
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import cm
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import (
        BaseDocTemplate, Frame, Image, KeepTogether, ListFlowable, ListItem, NextPageTemplate,
        PageBreak, PageTemplate, Paragraph, Preformatted, Spacer, Table, TableStyle,
    )
    from reportlab.platypus.tableofcontents import TableOfContents
    from xml.sax.saxutils import escape

    F = "C:/Windows/Fonts/"
    pdfmetrics.registerFont(TTFont("Body", F + "segoeui.ttf"))
    pdfmetrics.registerFont(TTFont("Body-Bold", F + "segoeuib.ttf"))
    pdfmetrics.registerFont(TTFont("Body-Italic", F + "segoeuii.ttf"))
    pdfmetrics.registerFont(TTFont("Body-BoldItalic", F + "segoeuiz.ttf"))
    pdfmetrics.registerFontFamily("Body", normal="Body", bold="Body-Bold", italic="Body-Italic", boldItalic="Body-BoldItalic")
    pdfmetrics.registerFont(TTFont("Mono", F + "consola.ttf"))
    pdfmetrics.registerFont(TTFont("Mono-Bold", F + "consolab.ttf"))
    pdfmetrics.registerFontFamily("Mono", normal="Mono", bold="Mono-Bold", italic="Mono", boldItalic="Mono-Bold")
    pdfmetrics.registerFont(TTFont("Sym", F + "seguisym.ttf"))
    body_cmap = FTFont(F + "segoeui.ttf").getBestCmap()
    mono_cmap = FTFont(F + "consola.ttf").getBestCmap()

    def with_fallback(text: str, cmap: dict, size: float) -> str:
        """Wrap characters the font lacks in a Segoe UI Symbol span."""
        out: list[str] = []
        buf: list[str] = []
        sym = False
        for ch in text:
            need = ord(ch) > 127 and ord(ch) not in cmap
            if need != sym:
                seg = "".join(buf)
                out.append(f'<font name="Sym" size="{size}">{seg}</font>' if sym else seg)
                buf = []
                sym = need
            buf.append(ch)
        seg = "".join(buf)
        out.append(f'<font name="Sym" size="{size}">{seg}</font>' if sym else seg)
        return "".join(out)

    def markup(text: str, size: float) -> str:
        parts: list[str] = []
        for style, chunk in inline_runs(text):
            e = with_fallback(escape(chunk), body_cmap, size)
            if style == "bold":
                parts.append(f"<b>{e}</b>")
            elif style == "italic":
                parts.append(f"<i>{e}</i>")
            elif style == "code":
                parts.append(f'<font name="Mono" size="{size - 0.5}">{e}</font>')
            else:
                parts.append(e)
        return "".join(parts)

    GREEN = colors.HexColor("#143D2E")
    styles = {
        "title": ParagraphStyle("title", fontName="Body-Bold", fontSize=38, leading=44, alignment=TA_CENTER, textColor=GREEN),
        "sub": ParagraphStyle("sub", fontName="Body", fontSize=18, leading=24, alignment=TA_CENTER),
        "sub2": ParagraphStyle("sub2", fontName="Body", fontSize=13, leading=18, alignment=TA_CENTER),
        "h1": ParagraphStyle("h1", fontName="Body-Bold", fontSize=17, leading=22, textColor=GREEN, spaceBefore=6, spaceAfter=8),
        "h2": ParagraphStyle("h2", fontName="Body-Bold", fontSize=13, leading=17, textColor=GREEN, spaceBefore=10, spaceAfter=5),
        "h3": ParagraphStyle("h3", fontName="Body-Bold", fontSize=11, leading=14, spaceBefore=8, spaceAfter=4),
        "body": ParagraphStyle("body", fontName="Body", fontSize=9.6, leading=13, spaceAfter=5),
        "quote": ParagraphStyle("quote", fontName="Body-Italic", fontSize=9.6, leading=13, leftIndent=18, rightIndent=18, spaceAfter=6, textColor=colors.HexColor("#2F4A3E")),
        "cell": ParagraphStyle("cell", fontName="Body", fontSize=7.6, leading=9.4),
        "cellh": ParagraphStyle("cellh", fontName="Body-Bold", fontSize=7.6, leading=9.4),
        "caption": ParagraphStyle("caption", fontName="Body-Italic", fontSize=8.5, leading=11, alignment=TA_CENTER, spaceBefore=3, spaceAfter=8),
        "toc0": ParagraphStyle("toc0", fontName="Body-Bold", fontSize=10.5, leading=15),
        "toc1": ParagraphStyle("toc1", fontName="Body", fontSize=9.5, leading=13, leftIndent=14),
    }

    PW, PH = A4
    margin = 1.8 * cm

    class Doc(BaseDocTemplate):
        def afterFlowable(self, fl):
            if isinstance(fl, Paragraph) and fl.style.name in ("h1", "h2"):
                level = 0 if fl.style.name == "h1" else 1
                text = fl.getPlainText()
                key = "toc-" + re.sub(r"\W+", "-", text)[:60]
                self.canv.bookmarkPage(key)
                self.canv.addOutlineEntry(text, key, level=level, closed=False)
                self.notify("TOCEntry", (level, text, self.page, key))

    doc = Doc(str(OUT_PDF), pagesize=A4, leftMargin=margin, rightMargin=margin, topMargin=margin, bottomMargin=margin,
              title="RASTA AI — SIH 2026 Day 1 Task 1 — System Design", author="NER-AI LOGISTICS (Team 17)")

    def on_page(canv, d):
        canv.saveState()
        w, h = canv._pagesize
        canv.setFont("Body", 7.5)
        canv.setFillColor(colors.HexColor("#555555"))
        canv.drawString(margin, 0.9 * cm, "RASTA AI · SIH 2026 · Day 1 Task 1 · Problem Understanding, Solution Planning & System Design")
        canv.drawRightString(w - margin, 0.9 * cm, f"Page {d.page}")
        canv.restoreState()

    portrait = PageTemplate("portrait", [Frame(margin, margin, PW - 2 * margin, PH - 2 * margin, id="f")], onPage=on_page)
    landscape_pt = PageTemplate("landscape", [Frame(margin, margin, PH - 2 * margin, PW - 2 * margin, id="f")], onPage=on_page, pagesize=LS(A4))
    doc.addPageTemplates([portrait, landscape_pt])

    story: list = []
    # ---- title page
    story += [Spacer(1, 5.5 * cm), Paragraph("RASTA AI", styles["title"]), Spacer(1, 0.4 * cm),
              Paragraph("SIH 2026 — Day 1 Task 1", styles["sub"]),
              Paragraph("Problem Understanding, Solution Planning &amp; System Design", styles["sub2"]), Spacer(1, 1.2 * cm)]
    meta = next((b for b in front if b.kind == "table"), None)
    if meta:
        rows = [[Paragraph(markup(r[0], 8.5), styles["cellh"]), Paragraph(markup(r[1] if len(r) > 1 else "", 8.5), styles["cell"])] for r in meta.rows if any(r)]
        t = Table(rows, colWidths=[4.5 * cm, 11.5 * cm])
        t.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#B8C4BE")), ("VALIGN", (0, 0), (-1, -1), "TOP"),
                               ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#EEF3F0")), ("LEFTPADDING", (0, 0), (-1, -1), 5), ("RIGHTPADDING", (0, 0), (-1, -1), 5)]))
        story.append(t)
    note = next((b for b in front if b.kind == "para" and b.text.startswith("**How to read")), None)
    if note:
        story += [Spacer(1, 0.8 * cm), Paragraph(markup(note.text, 9), styles["body"])]
    story.append(PageBreak())
    # ---- table of contents
    story.append(Paragraph("Table of contents", styles["h1"]))
    toc = TableOfContents()
    toc.levelStyles = [styles["toc0"], styles["toc1"]]
    toc.dotsMinLevel = 0
    story += [toc, PageBreak()]

    def para_style_for(level: int) -> ParagraphStyle:
        return {1: styles["h1"], 2: styles["h1"], 3: styles["h2"], 4: styles["h3"]}[level]

    def make_table(b: Block, avail: float) -> Table:
        cols = len(b.rows[0])
        # column widths proportional to the longest cell text, clamped
        lens = [max(len(r[c]) if c < len(r) else 0 for r in b.rows) for c in range(cols)]
        lens = [min(max(l, 16), 90) for l in lens]
        total = sum(lens)
        widths = [avail * l / total for l in lens]
        data = []
        for n, r in enumerate(b.rows):
            cells = []
            for c in range(cols):
                txt = r[c] if c < len(r) else ""
                cells.append(Paragraph(markup(txt, 7.6), styles["cellh"] if n == 0 else styles["cell"]))
            data.append(cells)
        t = Table(data, colWidths=widths, repeatRows=1)
        t.setStyle(TableStyle([
            ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#B8C4BE")),
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E3EBE6")),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("LEFTPADDING", (0, 0), (-1, -1), 4), ("RIGHTPADDING", (0, 0), (-1, -1), 4),
            ("TOPPADDING", (0, 0), (-1, -1), 3), ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
        ]))
        return t

    def code_flowable(b: Block, avail: float) -> Preformatted:
        lines = b.text.split("\n")
        longest = max((len(l) for l in lines), default=40)
        size = 8.0
        while size > 5.0 and longest * size * 0.55 > avail - 12:  # Consolas advance ≈ 0.55 em
            size -= 0.25
        while size > 5.5 and len(lines) * size * 1.22 > 660:  # keep a long diagram on one page
            size -= 0.25
        st = ParagraphStyle("code", fontName="Mono", fontSize=size, leading=size * 1.22, backColor=colors.HexColor("#F3F5F4"),
                            borderPadding=(4, 4, 4, 4), leftIndent=4, spaceAfter=8)
        return Preformatted(with_fallback_pre(b.text), st)

    def with_fallback_pre(text: str) -> str:
        # Preformatted has no inline markup; substitute the one glyph Consolas lacks.
        return text.replace("\u21c4", "<->")

    landscape = False
    avail_p = PW - 2 * margin
    avail_l = PH - 2 * margin
    first_h2 = True
    plan = landscape_plan(body)
    for b, want in zip(body, plan):
        just_broke = False
        if want and not landscape:
            story += [NextPageTemplate("landscape"), PageBreak()]
            landscape = True
            just_broke = True
        elif landscape and not want and b.kind in ("heading", "para", "bullets", "numbers", "code", "quote", "table", "image"):
            story += [NextPageTemplate("portrait"), PageBreak()]
            landscape = False
            just_broke = True
        avail = avail_l if landscape else avail_p
        if b.kind == "heading":
            if b.level == 2 and not first_h2 and not just_broke:
                story.append(PageBreak())
            first_h2 = False if b.level == 2 else first_h2
            story.append(Paragraph(markup(b.text, 14), para_style_for(b.level)))
        elif b.kind == "para":
            story.append(Paragraph(markup(b.text, 9.6), styles["body"]))
        elif b.kind == "quote":
            story.append(Paragraph(markup(b.text, 9.6), styles["quote"]))
        elif b.kind in ("bullets", "numbers"):
            items = [ListItem(Paragraph(markup(it, 9.6), styles["body"]), leftIndent=12) for it in b.items]
            story.append(ListFlowable(items, bulletType="bullet" if b.kind == "bullets" else "1", start=None if b.kind == "bullets" else 1,
                                      bulletFontName="Body", bulletFontSize=8, leftIndent=14))
        elif b.kind == "code":
            story.append(code_flowable(b, avail))
        elif b.kind == "table":
            story.append(make_table(b, avail))
            story.append(Spacer(1, 6))
        elif b.kind == "image":
            png = DIAGRAMS / f"diagram-{b.index}.png"
            if png.exists():
                from PIL import Image as PILImage
                iw, ih = PILImage.open(png).size
                maxw = avail
                maxh = (PW if landscape else PH) - 2 * margin - 2.2 * cm
                scale = min(maxw / iw, maxh / ih)
                img = Image(str(png), width=iw * scale, height=ih * scale)
                story.append(KeepTogether([img, Paragraph(DIAGRAM_CAPTIONS.get(b.index, f"Figure {b.index}"), styles["caption"])]))
        elif b.kind == "rule":
            pass
    doc.multiBuild(story)
    print("pdf:", OUT_PDF.name, OUT_PDF.stat().st_size // 1024, "KB")


if __name__ == "__main__":
    md = MD.read_text(encoding="utf8").replace("\r\n", "\n")
    blocks = parse(md)
    front, body = split_front_matter(blocks)
    print("blocks:", len(blocks), "front:", len(front), "body:", len(body))
    build_docx(front, body)
    build_pdf(front, body)
