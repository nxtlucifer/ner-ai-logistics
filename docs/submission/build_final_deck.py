"""Build the final SIH 2026 deck on the official template (6 slides).

    python docs/submission/build_final_deck.py
    -> docs/submission/RASTA_AI_SIH26002_TEAM17_FINAL.pptx

Every fact on the slides comes from docs/PPT_SOURCE_OF_TRUTH.md; every picture
is a real screen from docs/submission/screenshots/. The template's own pointers
(title fields, section titles, team oval, SIH logo, footer) are kept; only its
instruction text boxes and the decorative brain picture are replaced.
Rendering (PDF, PNG) is done afterwards with PowerPoint (render_final_deck.ps1).
"""
from __future__ import annotations

import os
import sys
from pathlib import Path

from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[2]
_WIN_TEMPLATE = Path(r"D:\SIH2026 PPT Format.pptx")
_REPO_TEMPLATE = Path(__file__).resolve().parent / "template.pptx"
TEMPLATE = (Path(os.environ["SIH_TEMPLATE"]) if os.environ.get("SIH_TEMPLATE")
            else (_WIN_TEMPLATE if _WIN_TEMPLATE.exists() else _REPO_TEMPLATE))
SHOTS = ROOT / "docs" / "submission" / "screenshots"
OUT = Path(os.environ.get("SIH_OUT", ROOT / "docs" / "submission" / "RASTA_AI_SIH26002_TEAM17_FINAL.pptx"))
CROPS = Path(os.environ.get("TEMP", ".")) / "sih" / "crops"
CROPS.mkdir(parents=True, exist_ok=True)

FONT = "Calibri"
INK = RGBColor(0x10, 0x18, 0x20)
BLUE = RGBColor(0x25, 0x63, 0xEB)
GREEN = RGBColor(0x05, 0x96, 0x69)
WARN = RGBColor(0xD9, 0x77, 0x06)
DANGER = RGBColor(0xDC, 0x26, 0x26)
GREY = RGBColor(0x6B, 0x72, 0x80)
MUTED = RGBColor(0x9C, 0xA3, 0xAF)
LINE = RGBColor(0xE5, 0xE7, 0xEB)
BG = RGBColor(0xF7, 0xF8, 0xFA)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)


# ---------------------------------------------------------------- helpers
def remove(shape):
    el = shape._element
    el.getparent().remove(el)


def no_shadow(shape):
    spPr = shape._element.spPr
    if spPr.find(qn("a:effectLst")) is None:
        etree.SubElement(spPr, qn("a:effectLst"))


def rect(slide, x, y, w, h, fill=None, line=None, rounded=False, line_w=0.75, radius=None):
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE
    s = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    if rounded:
        # adjustment 0 is a fraction of the SHORTER side, so a fixed fraction gives a
        # huge radius on a tall panel. radius= asks for a corner in inches instead.
        s.adjustments[0] = 0.12 if radius is None else min(0.5, radius / min(w, h))
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid(); s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line; s.line.width = Pt(line_w)
    no_shadow(s)
    s.text_frame.text = ""
    return s


def text(slide, x, y, w, h, paras, size=14, color=INK, bold=False, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, spacing=1.05, space_after=0, margin=0.02):
    """paras: list of paragraphs; a paragraph is a str or a list of (text, opts) runs.
    opts: bold, color, size, font, italic."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(margin)
    tf.margin_top = tf.margin_bottom = Inches(margin)
    first = True
    for para in paras:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = align
        p.line_spacing = spacing
        p.space_after = Pt(space_after)
        runs = [(para, {})] if isinstance(para, str) else para
        for t, o in runs:
            r = p.add_run(); r.text = t
            f = r.font
            f.name = o.get("font", FONT); f.size = Pt(o.get("size", size))
            f.bold = o.get("bold", bold); f.color.rgb = o.get("color", color)
            if o.get("italic"):
                f.italic = True
    return tb


def label(slide, x, y, w, t, color=GREY, size=10.5):
    return text(slide, x, y, w, 0.36, [[(t, {"bold": True, "color": color, "size": size})]])


def pill(slide, x, y, w, h, t, fill=INK, color=WHITE, size=10.5, line=None, bold=True):
    s = rect(slide, x, y, w, h, fill=fill, line=line, rounded=True)
    s.adjustments[0] = 0.5
    tf = s.text_frame
    tf.margin_left = tf.margin_right = Inches(0.06)
    tf.margin_top = tf.margin_bottom = Inches(0.0)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.word_wrap = True
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = t
    r.font.name = FONT; r.font.size = Pt(size); r.font.bold = bold; r.font.color.rgb = color
    return s


def box(slide, x, y, w, h, title, sub=None, fill=WHITE, line=LINE, title_color=INK,
        sub_color=GREY, title_size=11, sub_size=9, accent=None, align=PP_ALIGN.LEFT):
    s = rect(slide, x, y, w, h, fill=fill, line=line, rounded=True)
    if accent is not None:
        rect(slide, x, y + 0.08, 0.06, h - 0.16, fill=accent)
    paras = [[(title, {"bold": True, "color": title_color, "size": title_size})]]
    if sub:
        paras.append([(sub, {"color": sub_color, "size": sub_size})])
    text(slide, x + (0.14 if accent is not None else 0.1), y + 0.05, w - 0.2, h - 0.1, paras,
         anchor=MSO_ANCHOR.MIDDLE, spacing=1.0, align=align)
    return s


def panel(slide, x, y, w, h, title=None, sub=None, fill=WHITE, line=BLUE, line_w=1.25,
          title_size=11.5, title_color=None, align=PP_ALIGN.LEFT, pad=0.16):
    """A bordered, titled region — the SIH house convention: every block on the slide
    is framed and says what it is. Call before the content so it sits behind it.
    Returns the y at which the panel's own content may start."""
    rect(slide, x, y, w, h, fill=fill, line=line, rounded=True, line_w=line_w, radius=0.1)
    if not title:
        return y + pad
    runs = [(title, {"bold": True, "size": title_size, "color": title_color or BLUE})]
    if sub:
        runs.append((("   " + sub), {"size": title_size - 1.5, "color": GREY}))
    text(slide, x + pad, y + 0.08, w - 2 * pad, 0.3, [runs], align=align)
    return y + 0.42


def arrow(slide, x1, y1, x2, y2, color=MUTED, width=1.25, head=True):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = color; c.line.width = Pt(width)
    if head:
        ln = c.line._get_or_add_ln()
        tail = etree.SubElement(ln, qn("a:tailEnd")); tail.set("type", "triangle"); tail.set("w", "med"); tail.set("len", "med")
    return c


def hline(slide, x, y, w, color=LINE, width=0.75):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x), Inches(y), Inches(x + w), Inches(y))
    c.line.color.rgb = color; c.line.width = Pt(width)
    return c


def crop(name, boxpx=None, out=None):
    src = SHOTS / name
    im = Image.open(src).convert("RGB")
    if boxpx:
        im = im.crop(boxpx)
    dst = CROPS / (out or name)
    im.save(dst, "PNG", optimize=True)
    return dst, im.size


def picture(slide, path, x, y, w=None, h=None, size=None, frame=True):
    if size and w and not h:
        h = w * size[1] / size[0]
    if size and h and not w:
        w = h * size[0] / size[1]
    if frame:
        rect(slide, x - 0.03, y - 0.03, w + 0.06, h + 0.06, fill=WHITE, line=LINE, rounded=True)
    p = slide.shapes.add_picture(str(path), Inches(x), Inches(y), Inches(w), Inches(h))
    return p, w, h


def caption(slide, x, y, t, w=3.2):
    return pill(slide, x, y, w, 0.28, t, fill=INK, color=WHITE, size=9)


def vchain(slide, x, w, y, items, gap=0.13, h=0.3, size=10.5):
    """Vertical chain of pills; items: (text, fill, color, line). Returns the y after the last."""
    for i, (t, fill, color, line) in enumerate(items):
        pill(slide, x, y, w, h, t, fill=fill, color=color, line=line, size=size)
        if i < len(items) - 1:
            arrow(slide, x + w / 2, y + h + 0.01, x + w / 2, y + h + gap - 0.01, color=INK)
        y += h + gap
    return y


def team_oval(slide):
    for sh in slide.shapes:
        if sh.name.startswith("Oval"):
            sh.width = Inches(1.7); sh.height = Inches(0.8); sh.left = Inches(0.3); sh.top = Inches(0.3)
            tf = sh.text_frame; tf.word_wrap = True
            tf.margin_left = tf.margin_right = Inches(0.04)
            tf.vertical_anchor = MSO_ANCHOR.MIDDLE
            for p in list(tf.paragraphs)[1:]:
                p._p.getparent().remove(p._p)
            p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
            for r in list(p.runs):
                r._r.getparent().remove(r._r)
            r = p.add_run(); r.text = "NER-AI LOGISTICS"
            r.font.name = FONT; r.font.size = Pt(10); r.font.bold = True; r.font.color.rgb = INK
            return sh


def strip_instruction_box(slide):
    for sh in list(slide.shapes):
        if sh.name == "TextBox 8":
            remove(sh)


def set_title(slide, t, size=None):
    for sh in slide.shapes:
        if sh.is_placeholder and sh.placeholder_format.type in (1, 3):  # TITLE / CENTER_TITLE
            tf = sh.text_frame
            p = tf.paragraphs[0]
            for r in list(p.runs)[1:]:
                r._r.getparent().remove(r._r)
            if p.runs:
                p.runs[0].text = t
                if size:
                    p.runs[0].font.size = Pt(size)
                    p._p.get_or_add_endParaRPr().set("sz", str(int(size * 100)))
            else:
                tf.text = t
            return sh


def notes(slide, t):
    slide.notes_slide.notes_text_frame.text = t


# ---------------------------------------------------------------- slides
def slide1(s, nav, nav_size, mgr, mgr_size):
    for sh in list(s.shapes):
        if sh.name in ("Freeform: Shape 26", "Picture 4", "Subtitle 3", "TextBox 9"):
            remove(sh)
        elif sh.name == "Title 7":
            # The template's title box runs under the SIH logo, and at 40 pt the centred
            # "2026" prints on top of it in any renderer without Monotype Garamond.
            # Stop the box short of the logo and size the type to stay on one line.
            sh.width = Inches(10.15)
            for para in sh.text_frame.paragraphs:
                for r in para.runs:
                    r.font.size = Pt(34)
                para._p.get_or_add_endParaRPr().set("sz", "3400")
    # ---- left: who we are and what we were asked to solve
    panel(s, 0.45, 1.2, 6.85, 5.82)
    logo = ROOT / "driver-app" / "assets" / "brand-mark.png"
    s.shapes.add_picture(str(logo), Inches(0.68), Inches(1.4), Inches(0.82), Inches(0.82))
    text(s, 1.62, 1.36, 5.4, 0.62, [[("RASTA AI", {"bold": True, "size": 33, "color": INK})]], anchor=MSO_ANCHOR.MIDDLE)
    text(s, 1.64, 1.9, 5.5, 0.4, [[("Route Intelligence for Essential-Supply Logistics in North Eastern India", {"size": 12.5, "color": GREY})]])
    pill(s, 0.68, 2.42, 2.05, 0.6, "SIH26002", fill=BLUE, size=23)
    text(s, 2.92, 2.44, 4.2, 0.56, [[("Problem Statement ID", {"bold": True, "size": 10.5, "color": GREY})],
                                     [("MDoNER · Smart Automation · Software", {"size": 11.5, "color": INK})]], anchor=MSO_ANCHOR.MIDDLE)
    rows = [
        ("Problem Statement ID", "SIH26002"),
        ("Problem Statement Title", "AI-Based Smart Logistics and Accessibility Intelligence Platform for North Eastern Region (NER)"),
        ("Theme", "Smart Automation"),
        ("PS Category", "Software"),
        ("Team ID", "17  (internal group no.)"),
        ("Team Name", "NER-AI LOGISTICS"),
    ]
    y = 3.18
    for k, v in rows:
        tall = 0.48 if k == "Problem Statement Title" else 0.28
        text(s, 0.68, y, 2.15, tall, [[(k.upper(), {"bold": True, "size": 9.5, "color": GREY})]], anchor=MSO_ANCHOR.TOP)
        text(s, 2.83, y - 0.03, 4.28, tall, [[(v, {"bold": k in ("Problem Statement ID", "Team ID", "Team Name"), "size": 13, "color": INK})]], spacing=1.0)
        hline(s, 0.68, y + tall + 0.02, 6.4)
        y += tall + 0.06
    # the one line a judge must carry away from the title slide
    text(s, 0.68, 5.52, 6.4, 0.42, [[("SHORTEST IS NOT ALWAYS USABLE.", {"bold": True, "size": 20, "color": INK})]], spacing=1.0)
    text(s, 0.68, 5.97, 6.4, 0.38, [[("“Can this essential-supply truck reliably use this corridor now?”", {"size": 12.5, "color": BLUE})]], spacing=1.0)
    text(s, 0.68, 6.38, 6.4, 0.55, [[("End-to-end working prototype · physical-phone validated · live demo corridor Guwahati → Shillong", {"size": 10.5, "color": GREEN, "bold": True})]], spacing=1.1)
    # ---- right: the product itself, filling the panel rather than floating in it
    panel(s, 7.5, 1.2, 5.4, 5.82, "LIVE WORKING SYSTEM", sub="manager console + Android driver app")
    picture(s, mgr, 7.68, 2.46, h=3.88, size=mgr_size)
    picture(s, nav, 10.5, 1.72, h=4.62, size=nav_size)
    caption(s, 7.66, 6.48, "Manager route review · Driver navigation — real screens, real road", w=5.08)


def slide2(s, evid, evid_size):
    strip_instruction_box(s); team_oval(s)
    t = set_title(s, "RASTA AI — ROUTE INTELLIGENCE", size=28)
    t.left = Inches(2.1); t.width = Inches(8.5)
    label(s, 0.5, 1.2, 4, "PROPOSED SOLUTION")
    text(s, 0.5, 1.40, 8.7, 1.0, [[("NOT THE SHORTEST ROAD. ", {"bold": True, "size": 24, "color": INK}),
                                   ("THE ROAD THAT IS OPERATIONALLY USABLE NOW.", {"bold": True, "size": 24, "color": BLUE})]], spacing=1.0)
    text(s, 0.5, 2.28, 8.82, 0.55, [[("NER corridors: mountain terrain · monsoon · landslides · floods · weak signal · stale evidence. ", {"size": 11.5, "color": GREY}),
                                    ("A shortest-path router still sends the truck.", {"size": 11.5, "color": DANGER, "bold": True})]], spacing=1.0)
    # a normal router vs RASTA — each side framed and named, SIH house convention
    panel(s, 0.45, 2.88, 2.62, 3.4, "A NORMAL ROUTER", line=MUTED, title_color=GREY)
    y = vchain(s, 0.6, 2.32, 3.36, [("Origin", WHITE, INK, MUTED), ("Shortest / fastest road", WHITE, INK, MUTED), ("Truck", WHITE, INK, MUTED)], gap=0.26, h=0.3)
    text(s, 0.6, y + 0.06, 2.32, 1.3, [[("No terrain, weather, hazard or freshness check.", {"size": 10.5, "color": GREY})],
                                       [("No governance. No answer when data is missing.", {"size": 10.5, "color": GREY})]], spacing=1.05, space_after=4)
    panel(s, 3.2, 2.88, 6.12, 3.4, "RASTA AI", sub="evidence before the road")
    x, w = 3.35, 5.82
    y = vchain(s, x, w, 3.3, [("Origin", WHITE, INK, INK), ("Real route  ·  OSRM road, alternatives, turn steps", WHITE, INK, INK)], gap=0.1, h=0.28)
    # evidence block
    bh = 0.92
    rect(s, x, y, w, bh, fill=BG, line=LINE, rounded=True)
    text(s, x + 0.12, y + 0.04, w - 0.2, 0.26, [[("ROUTE-SPECIFIC EVIDENCE  ", {"bold": True, "size": 10.5, "color": INK}), ("sampled along the selected road, each factor with its freshness", {"size": 9.5, "color": GREY})]])
    chips = ["Terrain", "Weather", "Landslide exposure", "Flood context", "Official warnings", "Fleet traffic", "Evidence freshness"]
    cw = [0.85, 0.85, 1.35, 1.1, 1.25, 1.0, 1.35]
    cx, cy = x + 0.12, y + 0.32
    for c, wdt in zip(chips, cw):
        if cx + wdt > x + w - 0.1:
            cx = x + 0.12; cy += 0.28
        pill(s, cx, cy, wdt, 0.24, c, fill=WHITE, color=INK, line=INK, size=9, bold=False)
        cx += wdt + 0.08
    y += bh
    arrow(s, x + w / 2, y + 0.01, x + w / 2, y + 0.09, color=INK)
    y += 0.1
    y = vchain(s, x, w, y, [("OPERATIONAL DECISION  ·  CONTINUE / CAUTION / HOLD / REROUTE", BLUE, WHITE, None),
                            ("Manager governance  ·  reviews evidence, authorises, dispatches, approves reroutes", WHITE, INK, INK),
                            ("Driver  ·  navigation, danger context, 23 languages, offline package", WHITE, INK, INK)], gap=0.1, h=0.28)
    # unknown != safe — the punchline gets its own framed band
    panel(s, 0.45, 6.36, 8.87, 0.6, line=DANGER)
    pill(s, 0.6, 6.46, 2.3, 0.4, "UNKNOWN  ≠  SAFE", fill=DANGER, size=12.5)
    text(s, 3.02, 6.43, 6.18, 0.48, [[("The system refuses to fabricate certainty: ", {"size": 10.5, "color": INK, "bold": True}),
                                    ("missing or stale evidence is marked UNKNOWN, never counted as safe, and the route goes to an authorised reviewer.", {"size": 10.5, "color": INK})]], spacing=1.0)
    panel(s, 9.42, 1.2, 3.46, 5.74, "MANAGER · CHECK CONDITIONS", title_size=10.5)
    p, w, h = picture(s, evid, 9.64, 1.66, h=4.62, size=evid_size)
    caption(s, 9.57, 6.48, "Real screen · evidence · freshness · UNKNOWN", w=3.16)


def slide3(s):
    strip_instruction_box(s); team_oval(s)
    panel(s, 0.38, 1.12, 8.88, 3.94, "VERIFIED DATA → DECISION → DRIVER", sub="provider adapters with health + freshness")
    srcs = [("WEATHER", "Open-Meteo · MET Norway"), ("TERRAIN", "Copernicus DEM · OpenTopoData"), ("WARNINGS", "NDMA SACHET (CAP)"),
            ("FLOOD", "GloFAS"), ("LANDSLIDE HISTORY", "NASA GLC 2007–17"), ("TRAFFIC", "fleet GPS fixes")]
    x = 0.45; bw = 1.34; gap = 0.1; y = 1.58
    for t, sub in srcs:
        box(s, x, y, bw, 0.64, t, sub, title_size=10, sub_size=8.5)
        arrow(s, x + bw / 2, y + 0.66, x + bw / 2, y + 0.8, color=MUTED)
        x += bw + gap
    total = 6 * bw + 5 * gap
    box(s, 0.45, 2.4, total, 0.56, "ROUTE-SPECIFIC EVIDENCE  ·  FastAPI backend", "every factor sampled along the selected road · PostgreSQL + PostGIS (Supabase) · OSRM routes, alternatives, turn steps · freshness per factor",
        fill=BG, title_size=11, sub_size=9.5)
    arrow(s, 0.45 + total / 2, 2.98, 0.45 + total / 2, 3.1)
    box(s, 0.45, 3.12, total, 0.56, "DETERMINISTIC SAFETY POLICY", "11 factors · reason codes · UNKNOWN ≠ SAFE · reviewer authorisation when hazard data is missing",
        fill=BLUE, line=BLUE, title_color=WHITE, sub_color=WHITE, title_size=11.5, sub_size=9.5)
    arrow(s, 0.45 + total / 2, 3.7, 0.45 + total / 2, 3.82)
    dec = [("CONTINUE", GREEN), ("CAUTION", WARN), ("HOLD", DANGER), ("REROUTE", INK)]
    dw = (total - 3 * 0.15) / 4
    x = 0.45
    for t, c in dec:
        pill(s, x, 3.84, dw, 0.36, t, fill=c, size=11)
        x += dw + 0.15
    arrow(s, 0.45 + total / 2, 4.22, 0.45 + total / 2, 4.34)
    # the operating loop
    loop = [("MANAGER REVIEW", 1.45), ("DRIVER NAVIGATION", 1.6), ("ROUTE-AHEAD MONITOR (60 s)", 2.0), ("CONDITIONS CHANGE?", 1.55), ("REASSESS / REROUTE", 1.5)]
    x = 0.45; ly = 4.36
    for i, (t, w) in enumerate(loop):
        pill(s, x, ly, w, 0.36, t, fill=WHITE if i else INK, color=INK if i else WHITE, line=INK, size=9)
        if i < len(loop) - 1:
            arrow(s, x + w + 0.02, ly + 0.19, x + w + 0.1, ly + 0.19, color=INK)
        x += w + 0.11
    # return arrow: reassess -> policy
    rx = 0.45 + total + 0.12
    arrow(s, x - 0.11 + 0.02, ly + 0.18, rx, ly + 0.18, color=BLUE, head=False)
    arrow(s, rx, ly + 0.18, rx, 3.4, color=BLUE, head=False)
    arrow(s, rx, 3.4, 0.45 + total + 0.02, 3.4, color=BLUE)
    text(s, 0.45, 5.06, total, 0.36, [[("Decision loop: conditions are re-scored every 60 s along the road ahead; a material change goes back through the policy, and a reroute is proposed to the manager — never applied silently.", {"size": 9.5, "color": GREY})]])
    panel(s, 0.38, 5.46, 8.88, 1.38, "BUILT WITH")
    text(s, 0.54, 5.88, total - 0.2, 0.92, [
        [("CLIENTS  ", {"bold": True, "size": 10, "color": GREY}), ("Manager console — React + TypeScript, manager accounts only · Driver app — React Native / Expo, Android; one login, server-decided role, offline trip package", {"size": 10.5, "color": INK})],
        [("STACK  ", {"bold": True, "size": 10, "color": GREY}), ("FastAPI · PostgreSQL / PostGIS (Supabase) · React · React Native / Expo · OSRM · OpenStreetMap · Gemini", {"size": 10.5, "color": INK})],
        [("EVIDENCE  ", {"bold": True, "size": 10, "color": GREY}), ("Open-Meteo / MET Norway · NDMA SACHET · GloFAS · NASA historical landslides · OpenTopoData / Copernicus DEM", {"size": 10.5, "color": INK})],
    ], spacing=1.1, space_after=3)
    # right panel: supporting AI
    px = 9.42; pw = 3.46
    panel(s, px, 1.12, pw, 5.72, "WHO DECIDES", sub="AI is supporting", title_size=11)
    rows = [("DETERMINISTIC SAFETY ENGINE", "makes every operational decision", BLUE),
            ("GEMINI / OPENROUTER (LLM)", "explains the decision and assists in the driver's language — never decides", GREEN),
            ("EXPERIMENTAL ML", "research channel only; kept outside routing by the safety-validation gate", MUTED)]
    y = 1.62
    for t, sub, c in rows:
        rect(s, px + 0.15, y, 0.07, 0.74, fill=c)
        text(s, px + 0.3, y - 0.02, pw - 0.46, 0.82, [[(t, {"bold": True, "size": 11, "color": INK})], [(sub, {"size": 10.5, "color": GREY})]], spacing=1.0)
        y += 0.92
    hline(s, px + 0.15, 4.42, pw - 0.3, color=LINE)
    label(s, px + 0.15, 4.5, 3.1, "GOVERNANCE CHAIN")
    chain = ["Verified evidence", "Deterministic safety policy", "Manager review", "Driver action"]
    y = 4.84
    for i, t in enumerate(chain):
        pill(s, px + 0.15, y, pw - 0.3, 0.34, t, fill=WHITE, color=INK, line=INK, size=10.5)
        if i < 3:
            arrow(s, px + pw / 2, y + 0.36, px + pw / 2, y + 0.46, color=INK)
        y += 0.48


def slide4(s, mgr, mgr_size, truck, truck_size, nav, nav_size):
    strip_instruction_box(s); team_oval(s)
    text(s, 0.5, 1.06, 9.0, 0.5, [[("NOT A CONCEPT. A WORKING END-TO-END PROTOTYPE.", {"bold": True, "size": 22, "color": INK})]], anchor=MSO_ANCHOR.MIDDLE)
    steps = [("PLAN", 0.8), ("REVIEW", 0.95), ("DISPATCH", 1.1), ("ACCEPT", 1.0), ("TRUCK VERIFY", 1.35),
             ("NAVIGATE", 1.15), ("OFF-ROUTE", 1.15), ("MANAGER REROUTE APPROVAL", 2.25), ("DELIVERY", 1.1)]
    mgr_steps = (0, 1, 7)  # the manager governs these
    x = 0.5; sy = 1.62
    for i, (t, w) in enumerate(steps):
        pill(s, x, sy, w, 0.34, t, fill=BLUE if i in mgr_steps else (GREEN if i == 8 else WHITE),
             color=WHITE if i in mgr_steps or i == 8 else INK, line=None if i in mgr_steps or i == 8 else INK, size=9)
        if i < len(steps) - 1:
            arrow(s, x + w + 0.02, sy + 0.17, x + w + 0.09, sy + 0.17, color=INK)
        x += w + 0.11
    # three real screens — the hero of this slide, framed and named
    panel(s, 0.4, 2.12, 6.19, 4.3, "REAL SCREENS FROM THE CERTIFIED RUN", title_size=11)
    top = 2.58; hh = 3.44
    p, w1, h1 = picture(s, mgr, 0.56, top, h=hh, size=mgr_size)
    p, w2, h2 = picture(s, truck, 0.56 + w1 + 0.2, top, h=hh, size=truck_size)
    p, w3, h3 = picture(s, nav, 0.56 + w1 + w2 + 0.4, top, h=hh, size=nav_size)
    caption(s, 0.56, top + hh + 0.08, "Manager · route review", w=w1)
    caption(s, 0.56 + w1 + 0.2, top + hh + 0.08, "Driver · truck check", w=w2)
    caption(s, 0.56 + w1 + w2 + 0.4, top + hh + 0.08, "Driver · navigation", w=w3)
    # proof column — deliberately quieter than the screens on its left
    panel(s, 6.75, 2.12, 6.13, 4.3, "PROOF", sub="physical Android phone · hosted backend · real road", title_size=11)
    cx, cw = 6.91, 5.81
    text(s, cx, 2.62, cw, 0.62, [[("The whole lifecycle above was run end to end on a physical Android phone against the hosted backend, on the real Guwahati → Shillong road.", {"size": 11.5, "color": INK})]], spacing=1.05)
    y = 3.32
    for n, t in [("12/12", "physical judge-flow steps"), ("9/9", "physical role-flow steps")]:
        text(s, cx, y, 0.95, 0.34, [[(n, {"bold": True, "size": 15, "color": BLUE})]], anchor=MSO_ANCHOR.MIDDLE)
        text(s, cx + 1.0, y, cw - 1.0, 0.34, [[(t, {"size": 11, "color": INK})]], anchor=MSO_ANCHOR.MIDDLE)
        y += 0.4
    text(s, cx, y + 0.02, cw, 0.34, [[("1138+ ", {"bold": True, "size": 13, "color": BLUE}), ("backend  ·  ", {"size": 10.5, "color": INK}),
                                      ("621+ ", {"bold": True, "size": 13, "color": BLUE}), ("driver  ·  ", {"size": 10.5, "color": INK}),
                                      ("170+ ", {"bold": True, "size": 13, "color": BLUE}), ("manager tests", {"size": 10.5, "color": INK})]], anchor=MSO_ANCHOR.MIDDLE)
    y += 0.52
    pill(s, cx, y, cw, 0.38, "PHYSICAL ANDROID · CERTIFIED", fill=GREEN, size=11.5)
    text(s, cx, y + 0.48, cw, 0.36, [[("Canonical demo  ", {"bold": True, "size": 10.5, "color": GREY}), ("Guwahati → Shillong · ≈ 98.8 km real route", {"bold": True, "size": 11, "color": INK})]], spacing=1.0)
    text(s, cx, y + 0.86, cw, 0.4, [[("Every reroute on that run was approved by a manager before the phone followed it.", {"size": 10, "color": GREY})]], spacing=1.05)
    ly = 6.55
    hline(s, 0.5, ly - 0.05, 12.35)
    text(s, 0.5, ly, 12.35, 0.55, [[("KNOWN LIMITS → MITIGATION   ", {"bold": True, "size": 10, "color": GREY}),
                                   ("Render cold start → warm /health before the demo  ·  provider rate limits → cached evidence + explicit UNKNOWN  ·  background push needs Firebase config → in-app alerts already work  ·  experimental ML → outside routing until the gate passes", {"size": 10, "color": INK})]], spacing=1.05)
    notes(s, "Judge answer on ML: We trained and evaluated a landslide-hazard model, but deliberately keep it outside production routing because its geographic false-positive rate is still too high (recall 0.95 at FPR 0.50 on the NER holdout). This prevents unreliable ML from making safety-critical decisions. Tests: backend 1138 passed / 5 skipped, driver 621, manager 170 (14 Sep 2026, commit 0b89ddf). Physical certification: docs/terrain/HANDOFF.md section 9; judge flow 12/12, role flow 9/9 on APK 1.0.18.")


def slide5(s, lang, lang_size, mobile, mobile_size):
    strip_instruction_box(s); team_oval(s)
    text(s, 0.5, 1.12, 8.7, 0.9, [[("RASTA AI does not ask only ", {"size": 19, "color": GREY}), ("“Which road is shortest?”", {"size": 19, "color": INK, "bold": True}),
                                   ("  It asks ", {"size": 19, "color": GREY}), ("“Can this truck reliably use this corridor now?”", {"size": 19, "color": BLUE, "bold": True})]], spacing=1.05, anchor=MSO_ANCHOR.MIDDLE)
    text(s, 0.5, 2.02, 8.7, 0.34, [[("FOOD  ·  MEDICINE  ·  FUEL  ·  RELIEF SUPPLIES", {"bold": True, "size": 12.5, "color": GREEN}), ("   — essential logistics for hill communities, decided on evidence", {"size": 11.5, "color": GREY})]])
    panel(s, 0.42, 2.36, 8.95, 1.64, "WHO BENEFITS", title_size=11)
    who = [("FLEET MANAGERS", "Route usability and uncertainty visible before dispatch; auditable decisions."),
           ("DRIVERS", "Safer navigation, danger context, emergency support, 23 languages."),
           ("REMOTE COMMUNITIES", "More resilient access to food, medicine, fuel and relief supplies."),
           ("GOVERNMENT / OPERATIONS", "Auditable route evidence and disruption visibility per corridor.")]
    cw = 2.04
    for i, (h, b) in enumerate(who):
        x = 0.58 + i * (cw + 0.16)
        hline(s, x, 2.86, cw, color=BLUE, width=1.5)
        text(s, x, 2.92, cw, 1.02, [[(h, {"bold": True, "size": 11, "color": BLUE})], [(b, {"size": 11.5, "color": INK})]], spacing=1.05, space_after=3)
    panel(s, 0.42, 4.08, 8.95, 1.7, "DEMONSTRATED TODAY", title_size=11)
    caps = [("23", "selectable languages, explicit fallback status"), ("11", "route factors in every decision"), ("60 s", "route-ahead monitor while driving"),
            ("12/12", "judge-flow steps certified on a physical phone"), ("100 %", "of reroutes reviewed by a human"), ("0", "unknown factors counted as safe")]
    for i, (n, t) in enumerate(caps):
        col, row = divmod(i, 3)
        x = 0.62 + col * 4.35; y = 4.52 + row * 0.42
        text(s, x, y, 1.0, 0.4, [[(n, {"bold": True, "size": 18, "color": BLUE})]], anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.RIGHT)
        text(s, x + 1.08, y, 3.16, 0.4, [[(t, {"size": 11.5, "color": INK})]], anchor=MSO_ANCHOR.MIDDLE)
    panel(s, 0.42, 5.86, 8.95, 1.06, "HUMAN GOVERNANCE", sub="AI has no uncontrolled authority · no invented percentages", title_size=11)
    chain = ["Verified evidence", "Deterministic safety policy", "Manager governance", "Driver action"]
    x = 0.6; y = 6.34
    for i, t in enumerate(chain):
        w = 1.95
        pill(s, x, y, w, 0.38, t, fill=WHITE if i else INK, color=INK if i else WHITE, line=INK, size=10.5)
        if i < 3:
            arrow(s, x + w + 0.03, y + 0.19, x + w + 0.2, y + 0.19, color=INK)
        x += w + 0.26
    panel(s, 9.45, 1.2, 3.43, 5.72, "REAL SCREENS · APK 1.0.18", title_size=10.5)
    px = 9.61
    p, w, h = picture(s, lang, px, 1.68, h=2.5, size=lang_size)
    caption(s, px, 1.68 + h + 0.08, "Driver · 23 languages", w=w)
    p2, w2, h2 = picture(s, mobile, px + w + 0.16, 1.68, h=2.5, size=mobile_size)
    caption(s, px + w + 0.16, 1.68 + h2 + 0.08, "Manager · mobile", w=w2)
    text(s, px, 1.68 + h + 0.5, 3.11, 2.2, [[("Real screens, APK 1.0.18. One login, server-decided role: a manager gets a mobile fleet view, a driver gets navigation. Each language shows its status — Verified, Draft or English fallback — so nobody is misled.", {"size": 10, "color": GREY})]], spacing=1.05)


def slide6(s):
    strip_instruction_box(s); team_oval(s)
    # ---- two thirds of the slide: the sources the route decision actually rests on
    panel(s, 0.42, 1.12, 7.4, 4.9, "RESEARCH / EVIDENCE SOURCES BEHIND THE ROUTE DECISION", title_size=10.5)
    groups = [("ROUTING / PLACES", "OSRM · OpenStreetMap · Nominatim",
               "real road geometry, alternatives and turn steps; place search on the corridor"),
              ("WEATHER / TERRAIN", "Open-Meteo · MET Norway · Copernicus DEM · OpenTopoData",
               "forecast sampled along the road; climb, gradient and steep segments (GLO-90)"),
              ("OFFICIAL / FLOOD", "NDMA SACHET (CAP) · Copernicus GloFAS",
               "official warnings polled with freshness; river discharge against its 30-day mean"),
              ("LANDSLIDE HISTORY", "NASA Global Landslide Catalog",
               "historical exposure; inventory 2007–2017, ≤5 km location-accuracy filter"),
              ("AI ASSISTANCE", "Google Gemini",
               "explains the decision in the driver's language — never safety authority"),
              ("PROBLEM STATEMENT", "MDoNER · SIH26002",
               "AI-based smart logistics and accessibility intelligence for the NER")]
    gw = 3.4
    for i, (lab, names, what) in enumerate(groups):
        col, row = i % 2, i // 2
        x = 0.58 + col * (gw + 0.24)
        y = 1.62 + row * 1.44
        hline(s, x, y, gw, color=BLUE, width=1.5)
        text(s, x, y + 0.06, gw, 1.2,
             [[(lab, {"bold": True, "size": 10, "color": BLUE})],
              [(names, {"bold": True, "size": 11.5, "color": INK})],
              [(what, {"size": 10.5, "color": GREY})]], spacing=1.05, space_after=3)
    # ---- one third: the experimental model, and why it is not in the product
    mx = 8.25; mw = 4.6
    rect(s, mx - 0.18, 1.12, mw + 0.36, 4.9, fill=BG, line=BLUE, line_w=1.25, rounded=True, radius=0.1)
    label(s, mx, 1.26, mw, "EXPERIMENTAL LANDSLIDE RESEARCH")
    # 1. what the model achieved — large
    text(s, mx, 1.52, 1.85, 0.95, [[("95 %", {"bold": True, "size": 44, "color": INK})]], anchor=MSO_ANCHOR.MIDDLE)
    text(s, mx + 1.85, 1.56, mw - 1.85, 0.88,
         [[("RECALL", {"bold": True, "size": 12, "color": INK})],
          [("on the NER geographic holdout — 154 catalogued events, no NER row used in training", {"size": 9.5, "color": GREY})]],
         anchor=MSO_ANCHOR.MIDDLE, spacing=1.05)
    # 2. the decision that matters — dominant
    g = rect(s, mx, 2.62, mw, 0.82, fill=INK, rounded=True)
    text(s, mx, 2.68, mw, 0.72,
         [[("SAFETY GATE", {"bold": True, "size": 11, "color": MUTED})],
          [("NOT DEPLOYED", {"bold": True, "size": 21, "color": WHITE})]],
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, spacing=1.0)
    # 3. the reason it is gated — visible, honest, secondary
    text(s, mx, 3.52, mw, 0.36, [[("FPR 50 %", {"bold": True, "size": 13, "color": DANGER}),
                                  ("   too high for production deployment", {"size": 10, "color": GREY})]], anchor=MSO_ANCHOR.MIDDLE)
    flow = [[("TRAIN", 0.78), ("HELD-OUT NER TEST", 1.62), ("USEFUL SIGNAL FOUND", 1.78)],
            [("FALSE POSITIVES TOO HIGH", 1.98), ("SAFETY GATE", 1.08), ("NOT DEPLOYED", 1.26)]]
    fy = 3.94
    for r, row in enumerate(flow):
        x = mx
        for c, (t, w) in enumerate(row):
            last = (r == 1 and c == 2)
            pill(s, x, fy, w, 0.28, t, fill=INK if last else WHITE, color=WHITE if last else INK, line=INK, size=8)
            if not last:
                arrow(s, x + w + 0.01, fy + 0.14, x + w + 0.075, fy + 0.14, color=INK)
            x += w + 0.085
        fy += 0.4
    text(s, mx, 4.76, mw, 0.7, [[("VALIDATION BEFORE AUTOMATION. ", {"bold": True, "size": 11.5, "color": BLUE}),
                                 ("Strong recall was not enough. Held-out testing exposed the false-positive rate, so the gate keeps the model out of routing.", {"size": 10.5, "color": INK})]], spacing=1.05)
    text(s, mx, 5.46, mw, 0.5, [[("Precision 0.32 · F1 0.48 · PR-AUC 0.55 · ROC-AUC 0.83 · Brier 0.12 · calibrated logistic regression on NASA GLC + ERA5-Land rain + DEM slope · docs/HAZARD_ERROR_ANALYSIS.md", {"size": 8, "color": GREY})]], spacing=1.05)
    hline(s, 0.5, 6.16, 12.35)
    label(s, 0.5, 6.22, 3, "REFERENCES")
    refs = ["gpm.nasa.gov/landslides (NASA GLC)", "sachet.ndma.gov.in", "global-flood.emergency.copernicus.eu (GloFAS)", "open-meteo.com · api.met.no",
            "opentopodata.org · Copernicus DEM", "project-osrm.org · openstreetmap.org", "github.com/nxtlucifer/ner-ai-logistics (repository, tests, model registry, error analysis)"]
    text(s, 0.5, 6.45, 12.35, 0.6, [[("  ·  ".join(refs), {"size": 9.5, "color": GREY})]], spacing=1.1)
    notes(s, "Full references:\n"
             "NASA Global Landslide Catalog — https://gpm.nasa.gov/landslides/ (data via https://data.nasa.gov)\n"
             "NDMA SACHET CAP alerts — https://sachet.ndma.gov.in\n"
             "GloFAS (Copernicus Emergency Management Service) — https://global-flood.emergency.copernicus.eu ; flood API via https://open-meteo.com/en/docs/flood-api\n"
             "Open-Meteo forecast / ERA5-Land archive — https://open-meteo.com ; MET Norway Locationforecast — https://api.met.no\n"
             "OpenTopoData — https://www.opentopodata.org ; Copernicus DEM GLO-90 — https://spacedata.copernicus.eu\n"
             "OSRM — https://project-osrm.org ; OpenStreetMap — https://www.openstreetmap.org\n"
             "Repository — https://github.com/nxtlucifer/ner-ai-logistics (docs/MODEL_REGISTRY.md, docs/HAZARD_ERROR_ANALYSIS.md, docs/AI_INVENTORY.md)\n"
             "Model metrics (experimental, not deployed): NER holdout n=770, 154 events; recall 0.95, precision 0.32, F1 0.48, PR-AUC 0.55, ROC-AUC 0.83, FPR 0.50, Brier 0.12; temporal split recall 0.90 @ FPR 0.38.")


def main() -> int:
    prs = Presentation(str(TEMPLATE))
    nav, nav_size = crop("07-navigation.png")
    mgr, mgr_size = crop("02-trip-route.png", tuple(int(v) for v in os.environ.get("MGR_CROP", "1370,200,2470,1800").split(",")), "mgr-route.png")
    evid, evid_size = crop("03-route-evidence.png", tuple(int(v) for v in os.environ.get("EVID_CROP", "1400,70,2460,1720").split(",")), "mgr-evidence.png")
    truck, truck_size = crop("06-truck-verification.png", None if os.environ.get("TRUCK_WEB") else (0, 100, 1264, 2780), "truck-check.png")
    lang, lang_size = crop("11-language-selector.png", (0, 1000, 1264, 2780), "phone-lang.png")
    mobile, mobile_size = crop("10-manager-mobile.png", (0, 100, 1264, 2780), "phone-manager.png")
    slides = list(prs.slides)
    slide1(slides[0], nav, nav_size, mgr, mgr_size)
    slide2(slides[1], evid, evid_size)
    slide3(slides[2])
    slide4(slides[3], mgr, mgr_size, truck, truck_size, nav, nav_size)
    slide5(slides[4], lang, lang_size, mobile, mobile_size)
    slide6(slides[5])
    sldIdLst = prs.slides._sldIdLst
    last = list(sldIdLst)[6]
    prs.part.drop_rel(last.get(qn("r:id")))
    sldIdLst.remove(last)
    # the stock template ships as "Investor Pitch Deck Template" by "Crowdfunder";
    # that is what a judge's PDF viewer would put in the window title.
    cp = prs.core_properties
    cp.title = "RASTA AI — SIH26002 — Team NER-AI LOGISTICS (17)"
    cp.subject = "AI-Based Smart Logistics and Accessibility Intelligence Platform for North Eastern Region (NER)"
    cp.author = "NER-AI LOGISTICS"
    cp.last_modified_by = "NER-AI LOGISTICS"
    cp.category = "Smart India Hackathon 2026 — Smart Automation — Software"
    cp.keywords = "SIH26002, SIH 2026, RASTA AI, NER logistics, route intelligence"
    cp.comments = "Built by docs/submission/build_final_deck.py from docs/PPT_SOURCE_OF_TRUTH.md."
    prs.save(str(OUT))
    print("wrote", OUT, "slides:", len(prs.slides))
    return 0


if __name__ == "__main__":
    sys.exit(main())
